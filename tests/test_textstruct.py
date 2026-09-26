"""文本结构化层（app/textstruct）离线回归测试。

锁住三条最关键的性质：
1. **自足性** —— 每块都带章节路径前缀，脱离原文也能读懂；
2. **结构优先** —— 标题 / 代码块 / 表格是硬边界，不从中间切断；
3. **滑窗不丢路径** —— 块被切多片时，后续片必须补回前缀。
"""
from __future__ import annotations

import io

import pytest

from app.rag import chunk_blocks
from app.textparse import parse_document
from app.textstruct import clean_pdf_pages, parse_markdown, parse_plain_text, render_blocks


def _render(blocks) -> list[str]:
    return render_blocks(blocks, max_chars=600)


# --------------------------------------------------------------------------- #
# Markdown
# --------------------------------------------------------------------------- #
def test_markdown_blocks_carry_heading_path():
    md = "# 产品中心\n\n## 无线网桥\n\n### WB730\n\n点对点距离 5 公里。\n"
    out = _render(parse_markdown(md))
    assert any(b.startswith("【产品中心 > 无线网桥 > WB730】") for b in out)


def test_markdown_heading_pop_uses_level():
    """同级标题出现时，上一节的路径必须弹出，不能串成 A > B > C。"""
    md = "# 甲\n\n## 甲一\n\n内容一。\n\n# 乙\n\n内容二。\n"
    out = _render(parse_markdown(md))
    assert any(b.startswith("【甲 > 甲一】") for b in out)
    assert any(b.startswith("【乙】") for b in out)
    assert not any("甲" in b and "乙" in b for b in out)


def test_markdown_table_row_is_self_contained():
    """型号与它的参数必须在同一行记录里（不能整表糊成一段）。"""
    md = (
        "## 参数\n\n"
        "| 型号 | 重量 | 价格 |\n"
        "| --- | --- | --- |\n"
        "| WB730 | 15.5KG | 165 |\n"
        "| WB620E | 13KG | 145 |\n"
    )
    out = _render(parse_markdown(md))
    hit = [b for b in out if "WB730" in b]
    assert hit, "表格应产出记录块"
    row = [ln for ln in hit[0].split("\n") if "WB730" in ln]
    assert row, "每个型号应占一行"
    assert "15.5KG" in row[0] and "165" in row[0]


def test_markdown_short_rows_packed_but_long_rows_not():
    """短记录打包（避免几十字碎块淹没 top_k），长记录保持独立（避免主题稀释）。"""
    short_md = (
        "| 项 | 规则 |\n| --- | --- |\n"
        + "".join(f"| 规则{i} | 内容{i} |\n" for i in range(8))
    )
    packed = _render(parse_markdown(short_md, max_chars=600))
    assert len(packed) == 1, "8 行短记录应打包成一块"

    long_row = "本机型支持点对点传输，最大距离五公里，适用于园区与楼宇之间的桥接，支持 POE 供电。" * 3
    long_md = (
        "| 型号 | 说明 |\n| --- | --- |\n"
        f"| WB730 | {long_row} |\n"
        f"| WB620E | {long_row} |\n"
    )
    separate = _render(parse_markdown(long_md, max_chars=600))
    assert len(separate) == 2, "长记录必须各占一块"


def test_markdown_code_fence_is_atomic():
    md = "## 安装\n\n```bash\nmake install\nmake test\n```\n\n安装完成。\n"
    out = _render(parse_markdown(md))
    fenced = [b for b in out if "make install" in b]
    assert len(fenced) == 1
    assert "make test" in fenced[0]
    assert "```" in fenced[0]


# --------------------------------------------------------------------------- #
# 合并策略
# --------------------------------------------------------------------------- #
def test_merge_only_under_same_parent():
    blocks = parse_markdown(
        "# 章\n\n## 节一\n\n短内容一。\n\n## 节二\n\n短内容二。\n\n"
        "# 另一章\n\n另一章内容。\n"
    )
    out = _render(blocks)
    # 同父（章 > 节一 / 章 > 节二）的短块允许合并
    merged = [b for b in out if "短内容一" in b and "短内容二" in b]
    assert merged, "同父短块应合并"
    # 跨章绝不合并
    assert not any("短内容" in b and "另一章内容" in b for b in out)


# --------------------------------------------------------------------------- #
# 无格式文本
# --------------------------------------------------------------------------- #
def test_plain_text_numbered_heading():
    out = _render(parse_plain_text("第一章 总则\n\n本条款适用于所有供应商。\n\n第二章 细则\n\n细则内容。\n"))
    assert any(b.startswith("【第一章 总则】") for b in out)
    assert any(b.startswith("【第二章 细则】") for b in out)
    assert not any("总则" in b and "细则内容" in b for b in out), "不同章不能并一块"


def test_plain_text_list_item_is_not_heading():
    """'1. xxx' 这种单级编号是列表项，不能当标题。"""
    out = _render(parse_plain_text("1. 先把电源接上，确认指示灯亮起。\n2. 再按下配对键。\n"))
    assert not any(b.startswith("【") for b in out)


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #
def test_pdf_strips_repeated_header_footer():
    pages = [
        "说明书  第 1 页\n正文第一段。\n说明书  第 1 页",
        "说明书  第 2 页\n正文第二段。\n说明书  第 2 页",
        "说明书  第 3 页\n正文第三段。\n说明书  第 3 页",
    ]
    full, warns = clean_pdf_pages(pages)
    assert "说明书" not in full, "跨页重复的页眉/页脚必须剥掉"
    assert "正文第一段" in full and "正文第三段" in full
    assert any("页眉" in w for w in warns)


def test_pdf_rejoins_wrapped_lines():
    """被排版切断的行要拼回完整句子，否则每个块都是半截话。"""
    pages = [
        "本产品适用于硬质地面与短毛地毯的\n"
        "日常清洁，请勿用于长毛地毯。\n"
        "充电时请使用原装适配器。",
    ]
    full, _ = clean_pdf_pages(pages)
    assert "硬质地面与短毛地毯的日常清洁" in full


def test_pdf_heading_line_stays_separate():
    pages = ["第一章 安全须知\n请勿在潮湿环境中使用本\n产品，避免触电危险。"]
    full, _ = clean_pdf_pages(pages)
    lines = full.split("\n")
    assert lines[0] == "第一章 安全须知", "标题不能被并进正文"
    assert "使用本产品" in full


# --------------------------------------------------------------------------- #
# DOCX
# --------------------------------------------------------------------------- #
def test_docx_heading_and_table_in_document_order():
    docx = pytest.importorskip("docx")
    doc = docx.Document()
    doc.add_paragraph("产品手册", style="Heading 1")
    doc.add_paragraph("本章介绍安装。")
    doc.add_paragraph("安装准备", style="Heading 2")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "型号"
    table.cell(0, 1).text = "重量"
    table.cell(1, 0).text = "WB730"
    table.cell(1, 1).text = "15.5KG"
    doc.add_paragraph("安装步骤", style="Heading 2")
    doc.add_paragraph("第一步固定支架。")
    buf = io.BytesIO()
    doc.save(buf)

    res = parse_document("m.docx", buf.getvalue())
    joined = "\n".join(res.blocks)
    assert "【产品手册 > 安装准备】" in joined
    assert "【产品手册 > 安装步骤】" in joined
    # 表格挂在「安装准备」之后、「安装步骤」之前 —— 顺序不能乱
    assert joined.index("WB730") < joined.index("安装步骤")
    assert "15.5KG" in joined


# --------------------------------------------------------------------------- #
# 滑窗补前缀
# --------------------------------------------------------------------------- #
def test_sliding_window_keeps_path_prefix():
    para = "本机型支持点对点传输，最大距离为五公里，适用于园区桥接场景。" * 25
    res = parse_document("p.md", f"# 产品中心\n\n## 无线网桥\n\n{para}\n".encode("utf-8"))
    assert len(res.blocks) == 1
    pieces = chunk_blocks(res.blocks)
    assert len(pieces) >= 2, "长块应被滑窗切开"
    for p in pieces:
        assert p.startswith("【产品中心 > 无线网桥】"), "每一片都必须带路径前缀"


# --------------------------------------------------------------------------- #
# 图片：图本身进不了向量库，但绝不能静默丢弃
# --------------------------------------------------------------------------- #
def test_markdown_image_alt_is_kept():
    """![alt](url) 的 alt 是白捡的图描述，必须留下。"""
    out = _render(parse_markdown("# 手册\n\n![图 1-1 接线示意](a.png)\n\n正文。\n"))
    assert any("【图】图 1-1 接线示意" in b for b in out)


def test_markdown_image_without_alt_still_visible():
    """没有 alt 也要留占位，否则用户根本不知道文档里丢了一张图。"""
    out = _render(parse_markdown("# 手册\n\n![](a.png)\n\n正文。\n"))
    assert any("【图】" in b for b in out)


def test_docx_image_is_not_silently_dropped():
    """python-docx 只取 paragraph.text，图片段 text 为空 —— 以前直接 continue 静默丢图。"""
    import io

    import docx

    try:
        from PIL import Image
    except ImportError:  # pragma: no cover - 无 Pillow 时跳过
        return

    d = docx.Document()
    d.add_paragraph("产品手册", style="Heading 1")
    buf = io.BytesIO()
    Image.new("RGB", (8, 8), "red").save(buf, format="PNG")
    d.add_picture(io.BytesIO(buf.getvalue()))
    d.add_paragraph("图 2-3 设备安装示意图")

    res = parse_document("m.docx", _save(d))
    joined = "\n".join(res.blocks)
    assert "【图】图 2-3 设备安装示意图" in joined, "题注应当被认领为图描述"
    assert any("张图片" in w for w in res.warnings), "必须告知用户图未提取"


def test_docx_image_without_caption_uses_generic_placeholder():
    import io

    import docx

    try:
        from PIL import Image
    except ImportError:  # pragma: no cover
        return

    d = docx.Document()
    buf = io.BytesIO()
    Image.new("RGB", (8, 8), "red").save(buf, format="PNG")
    d.add_picture(io.BytesIO(buf.getvalue()))
    d.add_paragraph("后面是正常的段落。")

    res = parse_document("m.docx", _save(d))
    assert any("【图】" in b for b in res.blocks)


def _save(doc) -> bytes:
    import io

    b = io.BytesIO()
    doc.save(b)
    return b.getvalue()
