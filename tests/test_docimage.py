"""图片抽取与「图转文」回填的单元测试。全程离线，不调任何外部服务。"""
from __future__ import annotations

import io
import random
import struct
import zlib

import pytest

from app.docimage import (
    ImageRef,
    apply_descriptions,
    context_around,
    count_image_blocks,
    extract,
    image_size,
    mime_for,
)
from app.docstore import abs_path, remove_doc, save_image, save_original


def make_png(w: int = 120, h: int = 90) -> bytes:
    """手工构造 PNG（不依赖 Pillow）。

    填充必须用**随机噪声**而不是纯色：纯色 zlib 压缩后只有几百字节，
    会被自己的装饰性过滤（MIN_BYTES）拦掉，测不到正常路径。
    """
    rng = random.Random(20260924)

    def chunk(tag: bytes, data: bytes) -> bytes:
        return (
            struct.pack(">I", len(data)) + tag + data
            + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)
        )

    ihdr = struct.pack(">IIBBBBB", w, h, 8, 6, 0, 0, 0)
    raw = b"".join(b"\x00" + rng.randbytes(w * 4) for _ in range(h))
    return (
        b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr)
        + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")
    )


# --------------------------------------------------------------------------- #
# image_size
# --------------------------------------------------------------------------- #
def test_image_size_png():
    assert image_size(make_png(320, 200)) == (320, 200)


def test_image_size_jpeg():
    # JPEG 头没有真实图片，走到扫描循环应当安全返回 (0,0) 而不是抛异常
    assert image_size(b"\xff\xd8\xff\xe0" + b"\x00" * 20) == (0, 0)


def test_image_size_garbage():
    assert image_size(b"not an image") == (0, 0)


def test_mime_for():
    assert mime_for("png") == "image/png"
    assert mime_for(".jpg") == "image/jpeg"
    assert mime_for("webp") == "image/webp"


# --------------------------------------------------------------------------- #
# extract
# --------------------------------------------------------------------------- #
def test_extract_docx_images_in_order():
    from docx import Document

    doc = Document()
    doc.add_paragraph("第一章 安装说明")
    doc.add_picture(io.BytesIO(make_png(150, 120)))
    doc.add_paragraph("中间的正文")
    doc.add_picture(io.BytesIO(make_png(160, 130)))
    buf = io.BytesIO()
    doc.save(buf)

    result = extract(".docx", buf.getvalue())
    assert result.saw_image_marker
    assert len(result.images) == 2
    assert [r.seq for r in result.images] == [1, 2]
    assert all(r.mime == "image/png" for r in result.images)


def test_extract_skips_tiny_decorative():
    from docx import Document

    doc = Document()
    doc.add_picture(io.BytesIO(make_png(10, 8)))   # 尺寸过小
    buf = io.BytesIO()
    doc.save(buf)

    result = extract(".docx", buf.getvalue())
    assert result.saw_image_marker
    assert result.images == []
    assert any("装饰" in w for w in result.warnings)


def test_extract_dedupes_identical_bytes():
    from docx import Document

    same = make_png(140, 110)
    doc = Document()
    doc.add_picture(io.BytesIO(same))
    doc.add_paragraph("x")
    doc.add_picture(io.BytesIO(same))
    buf = io.BytesIO()
    doc.save(buf)

    result = extract(".docx", buf.getvalue())
    assert len(result.images) == 1, "同一张图重复引用只应识别一次"


def test_extract_unsupported_format_is_noop():
    result = extract(".txt", b"hello")
    assert result.images == []
    assert not result.saw_image_marker


def test_docx_hr_lines_are_not_images():
    """回归测试：Word 水平分隔线（空 v:imagedata，无关系引用）曾被当成 28 张
    假图，占位块污染知识库。判定收紧后必须 0 图。"""
    from docx import Document
    from docx.oxml import parse_xml

    from app.textstruct import parse_docx

    doc = Document()
    doc.add_paragraph("正文一段")
    hr = parse_xml(
        '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:v="urn:schemas-microsoft-com:vml" '
        'xmlns:o="urn:schemas-microsoft-com:office:office">'
        '<w:r><w:pict><v:rect o:hr="t"><v:imagedata o:title=""/></v:rect></w:pict></w:r></w:p>'
    )
    doc.element.body.append(hr)

    blocks, warnings = parse_docx(doc, max_chars=600)
    assert not any(b.body.startswith("【图】") for b in blocks), "分隔线不是图"
    assert not any("图片" in w for w in warnings)

    buf = io.BytesIO()
    doc.save(buf)
    result = extract(".docx", buf.getvalue())
    assert result.images == []
    assert not result.saw_image_marker


# --------------------------------------------------------------------------- #
# apply_descriptions
# --------------------------------------------------------------------------- #
def _ref(seq: int) -> ImageRef:
    return ImageRef(seq=seq, data=b"x", mime="image/png", ext="png", source="docx")


def test_apply_replaces_placeholders_in_order():
    blocks = ["正文一", "【图】图 1-1 接线示意", "正文二", "【图】图 1-2 面板布局", "正文三"]
    out, applied = apply_descriptions(blocks, [_ref(1), _ref(2)], {1: "端子A 接火线", 2: "指示灯 3 颗"})
    assert applied == 2
    assert out[1] == "【图】图 1-1 接线示意\n端子A 接火线"
    assert out[3] == "【图】图 1-2 面板布局\n指示灯 3 颗"
    assert out[0] == "正文一" and out[2] == "正文二"


def test_apply_keeps_placeholder_when_description_missing():
    blocks = ["【图】题注", "正文"]
    out, applied = apply_descriptions(blocks, [_ref(1)], {})
    assert applied == 0
    assert out[0] == "【图】题注", "没读出内容时必须保留原占位符，不能悄悄删掉"


def test_apply_falls_back_to_append_on_mismatch():
    """占位块数与图片数不一致时必须整体降级为追加，绝不就地错位配对。"""
    blocks = ["【图】A", "【图】B", "【图】C"]
    out, written = apply_descriptions(
        blocks, [_ref(1), _ref(2)], {1: "内容一", 2: "内容二"}
    )
    assert written == 2, "两张有描述的图都以独立块追加"
    assert out[:3] == blocks, "原有占位块原样保留（未就地改写）"
    assert out[3].startswith("【图】第 1 张插图\n内容一")
    assert out[4].startswith("【图】第 2 张插图\n内容二")


def test_apply_appends_for_pdf_with_page_number():
    ref = ImageRef(seq=1, data=b"x", mime="image/png", ext="png", source="pdf", page=3)
    out, applied = apply_descriptions(["正文"], [ref], {1: "参数表内容"})
    assert applied == 1
    assert out[1] == "【图】第 3 页插图\n参数表内容"


def test_count_image_blocks():
    assert count_image_blocks(["a", "【图】x", "b", "【图】y"]) == 2
    assert count_image_blocks([]) == 0


def test_context_around():
    blocks = ["前一节标题", "【图】图 1", "后一节正文"]
    ctx = context_around(blocks, 1)
    assert "前一节标题" in ctx
    assert "后一节正文" in ctx


def test_context_around_missing_seq():
    assert context_around(["正文"], 1) == ""


# --------------------------------------------------------------------------- #
# docstore
# --------------------------------------------------------------------------- #
def test_docstore_roundtrip(tmp_path, monkeypatch):
    monkeypatch.setattr("app.config.settings.assets_dir", str(tmp_path / "assets"))
    rel = save_original("t1", "d1", "手册.pdf", b"%PDF-1.4 fake")
    assert rel == "t1/d1/original.pdf"
    img_rel = save_image("t1", "d1", 1, "png", make_png())
    assert img_rel == "t1/d1/img_1.png"

    p = abs_path(img_rel)
    assert p is not None and p.is_file()


def test_docstore_rejects_path_escape(tmp_path, monkeypatch):
    monkeypatch.setattr("app.config.settings.assets_dir", str(tmp_path / "assets"))
    assert abs_path("../outside.png") is None
    assert abs_path("") is None


def test_docstore_remove(tmp_path, monkeypatch):
    monkeypatch.setattr("app.config.settings.assets_dir", str(tmp_path / "assets"))
    save_original("t1", "d1", "a.docx", b"x")
    remove_doc("t1", "d1")
    assert abs_path("t1/d1/original.docx") is None or not abs_path("t1/d1/original.docx").exists()


# --------------------------------------------------------------------------- #
# vision 文本清理（不发网络请求）
# --------------------------------------------------------------------------- #
def test_vision_clean_strips_fences_and_preamble():
    from app.vision import _clean

    assert _clean("```markdown\n| A | B |\n```") == "| A | B |"
    assert _clean("这张图展示了一台设备。正文") == "一台设备。正文", "客套开场白要剥掉"
    assert _clean("普通正文保持不变") == "普通正文保持不变"


def test_describe_many_returns_plain_text_dict(monkeypatch):
    """回归测试：describe_many 曾把 (文本, 错误) 元组当文本返回，
    上游直接写库炸出 sqlite3.ProgrammingError。返回值必须是 {seq: str}。"""
    import asyncio

    import app.vision as v

    async def fake_describe(data: bytes, mime: str, *, context: str = "", mode: str = "figure") -> str:
        return "识别出的文本"

    monkeypatch.setattr(v, "enabled", lambda: True)
    monkeypatch.setattr(v, "describe_image", fake_describe)
    out = asyncio.run(v.describe_many([(1, b"img", "image/png", ""), (2, b"img", "image/png", "")]))
    assert out == {1: "识别出的文本", 2: "识别出的文本"}
    assert all(isinstance(t, str) for t in out.values())


# --------------------------------------------------------------------------- #
# 上传链路集成测试（vision off：验证落盘 + 资产登记 + 提示，不碰外网）
# --------------------------------------------------------------------------- #
def test_upload_docx_with_image_registers_assets(client, star, tmp_path, monkeypatch):
    """带图 DOCX 上传后：原文件与图片落盘、资产行可查、占位说明出现在 warnings。"""
    from docx import Document

    from app.config import settings

    monkeypatch.setattr(settings, "vision_provider", "off")

    doc = Document()
    doc.add_paragraph("无线网桥参数说明")
    doc.add_picture(io.BytesIO(make_png(150, 120)))
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    res = client.post(
        "/api/knowledge-bases",
        json={"name": "图片链路验证库", "description": "临时"},
        headers=star,
    )
    assert res.status_code == 200, res.text
    kb_id = res.json()["id"]

    res = client.post(
        f"/api/knowledge-bases/{kb_id}/documents",
        files={
            "files": (
                "带图手册.docx",
                buf,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        headers=star,
    )
    assert res.status_code == 200, res.text
    item = res.json()["items"][0]
    assert item["status"] in ("success", "partial"), item
    assert any("图片" in w for w in item.get("warnings", []))

    doc_id = item["doc_id"]
    res = client.get(f"/api/documents/{doc_id}/assets", headers=star)
    assert res.status_code == 200
    assets = res.json()["items"]
    assert len(assets) == 1
    assert assets[0]["status"] == "skipped", "vision off 时应标记为 skipped"
    assert assets[0]["description"] == ""
    assert assets[0]["image_url"]

    res = client.get(assets[0]["image_url"], headers=star)
    assert res.status_code == 200
    assert res.content.startswith(b"\x89PNG\r\n\x1a\n"), "返回的是原图字节"


def test_upload_image_assets_isolated_between_tenants(client, star, sea, tmp_path, monkeypatch):
    """租户 B 看不到租户 A 的图片资产。"""
    from docx import Document

    from app.config import settings

    monkeypatch.setattr(settings, "vision_provider", "off")

    doc = Document()
    doc.add_picture(io.BytesIO(make_png(130, 100)))
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    res = client.post(
        "/api/knowledge-bases", json={"name": "隔离验证库", "description": ""}, headers=star
    )
    kb_id = res.json()["id"]
    res = client.post(
        f"/api/knowledge-bases/{kb_id}/documents",
        files={"files": ("a.docx", buf, "application/octet-stream")},
        headers=star,
    )
    doc_id = res.json()["items"][0]["doc_id"]

    assert client.get(f"/api/documents/{doc_id}/assets", headers=star).status_code == 200
    assert (
        client.get(f"/api/documents/{doc_id}/assets", headers=sea).status_code == 404
    ), "跨租户读取必须 404，不区分不存在与无权访问"


from app import docimage as _docimage_mod  # noqa: E402


# --------------------------------------------------------------------------- #
# 扫描版 PDF：整页渲染（页面级视觉转写的第一步）
# --------------------------------------------------------------------------- #
def _image_only_pdf(page_count: int = 2) -> bytes:
    """用 PyMuPDF 生成 N 页、每页贴一张纯色图的 PDF（等价扫描件）。"""
    pymupdf = pytest.importorskip("pymupdf")
    buf = io.BytesIO()
    with pymupdf.open() as doc:
        for _ in range(page_count):
            page = doc.new_page(width=300, height=400)
            pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 60, 60))
            pix.clear_with(200)
            page.insert_image(page.rect, pixmap=pix)
        doc.save(buf)
    return buf.getvalue()


def test_render_pdf_pages_returns_jpeg_per_page():
    data = _image_only_pdf(2)
    pages, warnings = _docimage_mod.render_pdf_pages(data, [1, 2], dpi=72)
    assert warnings == []
    assert [p for p, _, _ in pages] == [1, 2]
    for _, blob, mime in pages:
        assert mime == "image/jpeg"
        assert blob[:2] == b"\xff\xd8", "必须是 JPEG（体积可控，别把 PNG 整页送模型）"
        assert len(blob) > 1000


def test_render_pdf_pages_skips_out_of_range_with_warning():
    data = _image_only_pdf(1)
    pages, warnings = _docimage_mod.render_pdf_pages(data, [1, 9], dpi=72)
    assert [p for p, _, _ in pages] == [1]
    assert any("超出" in w for w in warnings)
